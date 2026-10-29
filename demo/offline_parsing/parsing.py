from __future__ import annotations

import math
import os
import re
import hashlib
import time
import unicodedata
from collections import Counter
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from contextlib import ExitStack
from itertools import islice
from pathlib import Path
from typing import Iterable, Iterator, TypeVar

from .models import ParsedDocument, ParsedPage, ParseOutcome


T = TypeVar("T")
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_PAGE_NUMBER_RE = re.compile(r"\d+")
_PAGE_MARGIN_RE = re.compile(
    r"^(?:(?:第\s*)?\d+\s*页?|[-—]\s*\d+\s*[-—]|page\s+\d+(?:\s+of\s+\d+)?)$",
    re.IGNORECASE,
)
_NOISE_LINE_RE = re.compile(r"^[\s\-_=*#~·•.。…]{4,}$")


def batched(values: Iterable[T], batch_size: int) -> Iterator[list[T]]:
    """按固定小批次惰性消费输入，避免先把全部文件放进内存。"""

    if batch_size <= 0:
        raise ValueError("batch_size must be greater than zero")
    iterator = iter(values)
    while batch := list(islice(iterator, batch_size)):
        yield batch


def partition_paths(paths: Iterable[Path]) -> tuple[list[Path], list[Path]]:
    """PDF 进入 CPU 进程池；Word/TXT 进入 IO 线程池。"""

    pdf_paths: list[Path] = []
    io_paths: list[Path] = []
    for path in paths:
        if path.suffix.lower() == ".pdf":
            pdf_paths.append(path)
        elif path.suffix.lower() in {".docx", ".txt"}:
            io_paths.append(path)
    return pdf_paths, io_paths


def recommended_worker_counts(cpu_count: int | None = None) -> tuple[int, int]:
    """CPU 池留一个核心给系统，IO 池可更大但封顶 32。"""

    cores = max(1, cpu_count or os.cpu_count() or 1)
    process_limit = 61 if os.name == "nt" else cores
    cpu_workers = min(max(1, cores - 1), process_limit)
    io_workers = max(4, min(32, cores * 4))
    return cpu_workers, io_workers


def clean_pages(
    pages: list[ParsedPage],
    *,
    repeated_margin_ratio: float = 0.6,
) -> list[ParsedPage]:
    """删除跨页重复的首尾模板行，再做保守文本规范化。

    这里没有 bbox 时只检查每页首行和尾行，不对全文做全局删除，避免误删正文
    中的同名标题或安全警告。生产版应在 bbox 顶部/底部候选区内执行同类规则。
    """

    if not 0 < repeated_margin_ratio <= 1:
        raise ValueError("repeated_margin_ratio must be in (0, 1]")
    if not pages:
        return []

    line_sets = [_normalized_nonempty_lines(page.text) for page in pages]
    first_counts = Counter(_margin_key(lines[0]) for lines in line_sets if lines)
    last_counts = Counter(_margin_key(lines[-1]) for lines in line_sets if lines)
    threshold = max(2, math.ceil(len(pages) * repeated_margin_ratio))
    repeated_first = {key for key, count in first_counts.items() if count >= threshold}
    repeated_last = {key for key, count in last_counts.items() if count >= threshold}

    cleaned_pages: list[ParsedPage] = []
    for page, lines in zip(pages, line_sets):
        if lines and _margin_key(lines[0]) in repeated_first:
            lines = lines[1:]
        if lines and _margin_key(lines[-1]) in repeated_last:
            lines = lines[:-1]
        text = clean_text("\n".join(lines))
        cleaned_pages.append(
            ParsedPage(
                page_number=page.page_number,
                text=text,
                width=page.width,
                height=page.height,
                metadata={**page.metadata, "cleaned": True},
            )
        )
    return cleaned_pages


def clean_text(text: str) -> str:
    """剔除替换字符、控制符、纯装饰行并压缩无意义重复符号。"""

    normalized = unicodedata.normalize("NFC", text).replace("\ufffd", "")
    normalized = _CONTROL_RE.sub("", normalized)
    normalized = re.sub(r"[!！]{3,}", "!!!", normalized)
    normalized = re.sub(r"[?？]{3,}", "???", normalized)
    normalized = re.sub(r"[ \t]+", " ", normalized)
    output_lines: list[str] = []
    blank_pending = False
    for raw_line in normalized.splitlines():
        line = raw_line.strip()
        if not line:
            blank_pending = bool(output_lines)
            continue
        if _NOISE_LINE_RE.fullmatch(line):
            continue
        if blank_pending:
            output_lines.append("")
            blank_pending = False
        output_lines.append(line)
    return "\n".join(output_lines).strip()


def discover_files(source_dir: Path) -> Iterator[Path]:
    if not source_dir.is_dir():
        raise NotADirectoryError(f"输入目录不存在或不是目录: {source_dir}")
    for path in sorted(source_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield path


def iter_parsed_documents(
    paths: Iterable[Path],
    *,
    source_dir: Path,
    parse_batch_size: int = 8,
    cpu_workers: int | None = None,
    io_workers: int | None = None,
) -> Iterator[ParseOutcome]:
    """分格式并行解析，并以小批次限制待处理 Future 与文件内存峰值。

    PDF 使用 ProcessPoolExecutor 隔离 CPU 密集任务；DOCX/TXT 使用
    ThreadPoolExecutor 等待磁盘 IO。每个文件由 _safe_parse 捕获异常，因此一个损坏
    文件不会终止整个批次。输出按 source 稳定排序，避免并发完成顺序改变 chunk ID。
    """

    if parse_batch_size <= 0:
        raise ValueError("parse_batch_size must be greater than zero")
    source_dir = source_dir.resolve()
    ordered_paths = sorted(
        (Path(path).resolve() for path in paths),
        key=lambda path: _relative_source(path, source_dir).casefold(),
    )
    pdf_paths, io_paths = partition_paths(ordered_paths)
    default_cpu, default_io = recommended_worker_counts()
    process_limit = 61 if os.name == "nt" else max(len(pdf_paths), 1)
    actual_cpu = min(
        max(1, cpu_workers or default_cpu),
        max(len(pdf_paths), 1),
        process_limit,
    )
    actual_io = min(max(1, io_workers or default_io), max(len(io_paths), 1))

    with ExitStack() as stack:
        pdf_pool = (
            stack.enter_context(ProcessPoolExecutor(max_workers=actual_cpu))
            if pdf_paths
            else None
        )
        io_pool = (
            stack.enter_context(ThreadPoolExecutor(max_workers=actual_io))
            if io_paths
            else None
        )
        for path_batch in batched(ordered_paths, parse_batch_size):
            futures = []
            for path in path_batch:
                executor = pdf_pool if path.suffix.lower() == ".pdf" else io_pool
                if executor is None:
                    continue
                futures.append(
                    executor.submit(_safe_parse, str(path), str(source_dir))
                )
            # submit 顺序来自规范化 source；result() 不会让完成顺序污染 ID。
            for future in futures:
                yield future.result()


def parse_document_file(path: Path, *, source_dir: Path) -> ParsedDocument:
    started = time.perf_counter()
    resolved = path.resolve(strict=True)
    suffix = resolved.suffix.lower()
    if suffix == ".pdf":
        pages, parser = _parse_pdf_file(resolved), "pymupdf"
    elif suffix == ".docx":
        pages, parser = _parse_docx_file(resolved), "python-docx"
    elif suffix == ".txt":
        pages, parser = _parse_txt_file(resolved), "text"
    else:
        raise ValueError(f"不支持的文件格式: {suffix}")
    cleaned_pages = clean_pages(pages)
    quality_status, quality_warnings = _assess_parse_quality(cleaned_pages, suffix)
    source = _relative_source(resolved, source_dir.resolve())
    document_id = _document_id(source, resolved)
    return ParsedDocument(
        document_id=document_id,
        source=source,
        source_path=str(resolved),
        pages=cleaned_pages,
        parser=parser,
        parse_seconds=time.perf_counter() - started,
        metadata={
            "source_hash": _streaming_sha256(resolved),
            "parser_version": _parser_version(parser),
            "cleaner_version": "demo-cleaner-v1",
            "parse_quality_status": quality_status,
            "parse_warnings": quality_warnings,
            "ocr_required": "ocr_required" in quality_warnings,
        },
        quality_status=quality_status,
        quality_warnings=quality_warnings,
    )


def iter_pdf_page_batches(
    path: Path,
    *,
    page_batch_size: int = 16,
    workers: int | None = None,
) -> Iterator[list[ParsedPage]]:
    """单个超大 PDF 的页范围多进程解析示例。

    每个子进程重新打开 PDF，只返回一个页批次。调用方应立刻切片/向量化并释放该批次。
    该低内存路径只做基础文本提取；跨页页眉页脚统计需要额外的两遍扫描。
    """

    if page_batch_size <= 0:
        raise ValueError("page_batch_size must be greater than zero")
    fitz = _import_fitz()
    with fitz.open(path) as document:
        page_count = document.page_count
    ranges = [
        (str(path.resolve()), start, min(start + page_batch_size, page_count))
        for start in range(0, page_count, page_batch_size)
    ]
    if not ranges:
        return
    default_workers, _ = recommended_worker_counts()
    process_limit = 61 if os.name == "nt" else len(ranges)
    actual_workers = min(
        max(1, workers or default_workers),
        len(ranges),
        process_limit,
    )
    with ProcessPoolExecutor(max_workers=actual_workers) as executor:
        for range_batch in batched(ranges, actual_workers * 2):
            for pages in executor.map(_parse_pdf_page_range, range_batch):
                yield pages


def _safe_parse(path_text: str, source_dir_text: str) -> ParseOutcome:
    path = Path(path_text)
    source_dir = Path(source_dir_text)
    source = _relative_source(path, source_dir)
    started = time.perf_counter()
    try:
        document = parse_document_file(path, source_dir=source_dir)
    except Exception as exc:  # noqa: BLE001 - Demo 需要按单文件隔离失败。
        return ParseOutcome(
            source=source,
            error_type=type(exc).__name__,
            error_message=str(exc),
            parse_seconds=time.perf_counter() - started,
        )
    return ParseOutcome(
        source=source,
        document=document,
        parse_seconds=document.parse_seconds,
    )


def _parse_txt_file(path: Path) -> list[ParsedPage]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    return [ParsedPage(page_number=1, text=text)]


def _parse_docx_file(path: Path) -> list[ParsedPage]:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise RuntimeError("缺少 python-docx，无法解析 DOCX") from exc
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append("\t".join(cell.text.strip() for cell in row.cells))
    return [ParsedPage(page_number=1, text="\n".join(paragraphs))]


def _parse_pdf_file(path: Path) -> list[ParsedPage]:
    fitz = _import_fitz()
    pages: list[ParsedPage] = []
    with fitz.open(path) as document:
        for page_index, page in enumerate(document):
            pages.append(
                ParsedPage(
                    page_number=page_index + 1,
                    text=page.get_text("text", sort=True),
                    width=float(page.rect.width),
                    height=float(page.rect.height),
                )
            )
    return pages


def _parse_pdf_page_range(task: tuple[str, int, int]) -> list[ParsedPage]:
    path_text, start, end = task
    fitz = _import_fitz()
    pages: list[ParsedPage] = []
    with fitz.open(path_text) as document:
        for page_index in range(start, end):
            page = document.load_page(page_index)
            pages.append(
                ParsedPage(
                    page_number=page_index + 1,
                    text=page.get_text("text", sort=True),
                    width=float(page.rect.width),
                    height=float(page.rect.height),
                )
            )
    return pages


def _import_fitz():
    try:
        import pymupdf

        return pymupdf
    except ModuleNotFoundError:
        try:
            import fitz

            return fitz
        except ModuleNotFoundError as exc:
            raise RuntimeError("缺少 PyMuPDF，无法解析 PDF") from exc


def _relative_source(path: Path, source_dir: Path) -> str:
    try:
        return path.resolve().relative_to(source_dir.resolve()).as_posix()
    except (ValueError, OSError):
        return path.name


def _streaming_sha256(path: Path, block_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while block := handle.read(block_size):
            digest.update(block)
    return digest.hexdigest()


def _document_id(source: str, path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(source.encode("utf-8"))
    digest.update(b"\0")
    with path.open("rb") as handle:
        while block := handle.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()[:24]


def _parser_version(parser: str) -> str:
    if parser == "text":
        return "stdlib-text-v1"
    if parser == "python-docx":
        try:
            import docx

            return f"python-docx-{docx.__version__}"
        except (ModuleNotFoundError, AttributeError):
            return "python-docx-unknown"
    fitz = _import_fitz()
    version = getattr(fitz, "version", None)
    if isinstance(version, tuple):
        version = version[0]
    return f"pymupdf-{version or 'unknown'}"


def _assess_parse_quality(
    pages: list[ParsedPage],
    suffix: str,
) -> tuple[str, list[str]]:
    text = "\n".join(page.text for page in pages).strip()
    if text:
        return "parsed_success", []
    warnings = ["empty_text"]
    if suffix == ".pdf":
        warnings.append("ocr_required")
    return "needs_review", warnings


def _normalized_nonempty_lines(text: str) -> list[str]:
    return [line.strip() for line in clean_text(text).splitlines() if line.strip()]


def _margin_key(line: str) -> str:
    normalized = line.strip().casefold()
    if _PAGE_MARGIN_RE.fullmatch(normalized):
        return _PAGE_NUMBER_RE.sub("#", normalized)
    return normalized
