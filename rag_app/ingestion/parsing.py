"""离线文档解析与清洗。"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import unicodedata
from dataclasses import dataclass
from html.parser import HTMLParser
from json import JSONDecodeError
from pathlib import Path
from typing import Any

from rag_app.core.error_codes import extract_error_codes, normalize_error_code
from rag_app.core.models import Document, ParsedBlock


PARSER_VERSION = "document-parser-v6"

TITLE_RE = re.compile(r"^#{1,6}\s+(.+?)\s*$", re.MULTILINE)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
FRONT_MATTER_RE = re.compile(r"\A---\n(.*?)\n---\n?", re.DOTALL)
JSON_RECORD_KEYS = ("records", "items", "data", "rows", "cases", "orders", "errors")
PDF_MIN_TEXT_CHARS_PER_PAGE = 30
PDF_GARBLED_RATIO_WARNING = 0.2
PDF_OCR_CONFIDENCE_WARNING = 0.75
PDF_TABLE_CELL_SPLIT_RE = re.compile(r"\t+|\s{2,}")
PAGE_TITLE_RE = re.compile(r"^第\s*(\d+)\s*页$")
SLIDE_TITLE_RE = re.compile(r"^第\s*(\d+)\s*页幻灯片$")
SEGMENT_TITLE_RE = re.compile(r"^片段\s+\d+\s+\[(.+)\]$")


@dataclass(frozen=True)
class ParsedFile:
    """文件解析后的正文和元数据。"""

    text: str
    metadata: dict


@dataclass(frozen=True)
class SubtitleCue:
    """字幕分段结果。"""

    start: str
    end: str
    text: str


@dataclass(frozen=True)
class PdfPageReport:
    """单页 PDF 文本抽取质量报告。"""

    page_number: int
    raw_text_length: int
    clean_text_length: int
    status: str
    warning_codes: list[str]
    table_like_lines: int = 0
    extraction_method: str = "text"
    ocr_confidence: float | None = None


@dataclass(frozen=True)
class PdfPageContent:
    """PDF 单页最终采用的文本内容。"""

    page_number: int
    raw_text: str
    text: str
    table_like_lines: int
    extraction_method: str = "text"
    ocr_confidence: float | None = None


@dataclass(frozen=True)
class PdfOcrPageResult:
    """OCR 后的单页文本和置信度。"""

    page_number: int
    text: str
    confidence: float | None = None


def parse_document_file(path: Path, source_dir: Path | None = None) -> Document | None:
    """解析单个知识文件，返回可进入离线切片流程的文档。"""

    # 统一入口：各种文件先转成 ParsedFile，再补齐可追溯元数据。
    # 后续切片、召回引用、灰度问题排查都依赖这些 metadata。
    parsed = _parse_file(path)
    if parsed is None or not parsed.text.strip(): #空文档不会进入解析流程
        return None
    #生成来源路径
    source_path = str(path.resolve())
    relative_source = _relative_source(path, source_dir)
    #计算内容哈希
    content_hash = hashlib.sha256(parsed.text.encode("utf-8")).hexdigest()

    document_id = _document_id(source_path, content_hash)
    version_id = _document_version_id(content_hash)
    #提取标题、来源类型、业务模块、租户和权限
    title = parsed.metadata.get("title") or _extract_title(parsed.text) or path.stem
    source_type = _infer_source_type(path, parsed.text)
    business_module = _infer_business_module(path, parsed.text)
    tenant_id = _normalize_tenant_id(parsed.metadata)
    permission_tags = _normalize_permission_tags(parsed.metadata)
    error_codes = extract_error_codes(parsed.text)
    #构造medatada数据
    metadata = {
        "source": relative_source,
        "source_path": source_path,
        "filename": path.name,
        "extension": path.suffix.lower(),
        "title": title,
        "source_type": source_type,
        "business_module": business_module,
        "content_hash": content_hash,
        "version_id": version_id,
        "tenant_id": tenant_id,
        "permission_tags": permission_tags,
        "parser_name": parsed.metadata.get("parser_name", "text"),
        "parser_version": PARSER_VERSION,
        "parse_quality_status": parsed.metadata.get("parse_quality_status", "parsed_success"),
        "parse_warnings": parsed.metadata.get("parse_warnings", []),
        "ocr_required": bool(parsed.metadata.get("ocr_required", False)),
        "error_codes": error_codes,
        **parsed.metadata, #如果 parsed.metadata 里也有同名字段，它会覆盖前面已经设置的字段。
    }
    metadata["parser_version"] = PARSER_VERSION
    metadata["tenant_id"] = tenant_id
    metadata["permission_tags"] = permission_tags
    metadata["parse_quality_status"] = metadata.get("parse_quality_status") or "parsed_success"
    metadata["parse_warnings"] = list(metadata.get("parse_warnings") or [])
    metadata["ocr_required"] = bool(metadata.get("ocr_required", False))
    metadata["error_codes"] = _normalize_error_code_list(
        metadata.get("error_codes") or error_codes
    )
    if len(metadata["error_codes"]) == 1:
        metadata["error_code"] = metadata["error_codes"][0]
    #完整文档构造为结构化block
    parsed_blocks = _build_parsed_blocks(
        document_id=document_id,
        version_id=version_id,
        text=parsed.text,
        document_metadata=metadata,
    )
    metadata["block_count"] = len(parsed_blocks)

    return Document(
        id=document_id,
        text=parsed.text,
        metadata=metadata,
        parsed_blocks=parsed_blocks,
    )


def clean_text(text: str) -> str:
    """清洗文本格式，保留 Markdown 结构和段落边界。"""

    # 不做激进清洗，避免破坏 Markdown 标题、列表和段落，这些结构会被切片器继续使用。
    text = text.replace("\ufeff", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    text = "\n".join(lines).strip()
    return re.sub(r"\n{3,}", "\n\n", text)


def _parse_file(path: Path) -> ParsedFile | None:
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        return _parse_html(path)
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix == ".json":
        return _parse_json(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".xlsx":
        return _parse_xlsx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".srt":
        return _parse_srt(path)
    if suffix == ".vtt":
        return _parse_vtt(path)
    return _parse_text(path)


def _parse_text(path: Path) -> ParsedFile | None:
    text = _read_text_with_fallback(path)
    front_matter, text = _extract_front_matter(text)
    cleaned = clean_text(text)
    if not cleaned:
        return None
    title = front_matter.get("title") or _extract_title(cleaned) or path.stem
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": "markdown" if path.suffix.lower() in {".md", ".rst"} else "text",
            "title": title,
            **front_matter,
        },
    )


def _parse_html(path: Path) -> ParsedFile | None:
    text = _read_text_with_fallback(path)
    if not text.strip():
        return None

    parser = _HtmlToMarkdownParser(default_title=path.stem)
    parser.feed(text)
    parser.close()
    cleaned = clean_text("\n".join(parser.markdown_lines()))
    if not cleaned or cleaned == f"# {parser.title or path.stem}":
        return None
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": "html",
            "title": parser.title or _extract_title(cleaned) or path.stem,
            "media_type": "html",
            "html_link_count": parser.link_count,
            "html_table_count": parser.table_count,
        },
    )


def _parse_docx(path: Path) -> ParsedFile | None:
    try:
        from docx import Document as DocxFile
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "解析 DOCX 需要安装可选依赖：pip install -e \".[office]\""
        ) from exc

    # Word 文档保留标题层级和表格行语义，避免制度、手册类文档被直接拍平成普通文本。
    docx = DocxFile(str(path))
    lines = [f"# {path.stem}"]
    paragraph_count = 0
    for paragraph in docx.paragraphs:
        content = clean_text(paragraph.text or "")
        if not content:
            continue
        paragraph_count += 1
        style_name = (paragraph.style.name if paragraph.style else "").lower()
        heading_level = _docx_heading_level(style_name)
        if heading_level is not None:
            lines.append("")
            lines.append(f"{'#' * min(heading_level + 1, 6)} {content}")
        else:
            lines.append(content)

    table_count = 0
    for table in docx.tables:
        table_count += 1
        lines.extend(["", f"## 表格 {table_count}"])
        table_lines = _table_rows_to_lines(
            rows=[
                [clean_text(cell.text or "") for cell in row.cells]
                for row in table.rows
            ]
        )
        lines.extend(table_lines)

    cleaned = clean_text("\n".join(lines))
    if cleaned == f"# {path.stem}":
        return None
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": "docx",
            "title": _extract_title(cleaned) or path.stem,
            "media_type": "docx",
            "paragraph_count": paragraph_count,
            "table_count": table_count,
        },
    )


def _parse_xlsx(path: Path) -> ParsedFile | None:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "解析 XLSX 需要安装可选依赖：pip install -e \".[office]\""
        ) from exc

    # Excel 以工作表和行级记录为主，保留列名=单元格值的对应关系，后续可做行级召回。
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    lines = [f"# {path.stem}"]
    sheet_count = 0
    row_count = 0
    try:
        for worksheet in workbook.worksheets:
            sheet_count += 1
            lines.extend(["", f"## 工作表：{worksheet.title}"])
            rows = [
                [
                    _stringify_excel_cell(value)
                    for value in row
                ]
                for row in worksheet.iter_rows(values_only=True)
            ]
            table_lines = _table_rows_to_lines(rows)
            row_count += max(len(rows) - 1, 0)
            lines.extend(table_lines)
    finally:
        workbook.close()

    cleaned = clean_text("\n".join(lines))
    if cleaned == f"# {path.stem}":
        return None
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": "xlsx",
            "title": path.stem,
            "media_type": "xlsx",
            "sheet_count": sheet_count,
            "row_count": row_count,
        },
    )


def _parse_csv(path: Path) -> ParsedFile | None:
    text = _read_text_with_fallback(path)
    if not text.strip():
        return None

    # CSV 不是直接拼成一坨文本，而是转成 Markdown 记录。
    # 这样每一行都有稳定的小节标题，后续召回时能定位到具体记录。
    rows = list(csv.reader(text.splitlines()))
    if not rows:
        return None

    headers = [cell.strip() for cell in rows[0]]
    data_rows = rows[1:]
    title = path.stem
    lines = [f"# {title}"]

    for index, row in enumerate(data_rows, start=1):
        if not any(cell.strip() for cell in row):
            continue
        lines.append("")
        lines.append(f"## 记录 {index}")
        for column_index, value in enumerate(row):
            value = value.strip()
            if not value:
                continue
            header = headers[column_index] if column_index < len(headers) and headers[column_index] else f"字段{column_index + 1}"
            lines.append(f"- {header}: {value}")

    cleaned = clean_text("\n".join(lines))
    if cleaned == f"# {title}":
        return None
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": "csv",
            "title": title,
            "row_count": len(data_rows),
            "columns": headers,
        },
    )


def _parse_json(path: Path) -> ParsedFile | None:
    text = _read_text_with_fallback(path)
    if not text.strip():
        return None

    # JSON 支持数组、单对象，以及常见业务包装字段。
    # 统一转成 Markdown 后，向量化阶段不用关心原始结构来自哪个字段。
    try:
        data = json.loads(text)
    except JSONDecodeError as exc:
        raise ValueError(f"JSON 文件解析失败: {path}，请检查格式是否正确。") from exc

    title, records = _json_records(path, data)
    lines = [f"# {title}"]
    for index, record in enumerate(records, start=1):
        lines.append("")
        lines.append(f"## {_json_record_title(index, record)}")
        if isinstance(record, dict):
            for key, value in record.items():
                if value in ("", None, [], {}):
                    continue
                lines.append(f"- {key}: {_stringify_json_value(value)}")
        else:
            lines.append(_stringify_json_value(record))

    cleaned = clean_text("\n".join(lines))
    if cleaned == f"# {title}":
        return None
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": "json",
            "title": title,
            "record_count": len(records),
            "json_root_type": type(data).__name__,
        },
    )


def _parse_pdf(path: Path) -> ParsedFile | None:
    try:
        import fitz
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "解析 PDF 需要安装 PyMuPDF：pip install -e \".[multimodal]\""
        ) from exc

    # PDF 先按文本型文档处理，只有空页、低文本页或疑似扫描页才触发 OCR。
    # OCR 是保底机制，不影响普通文字 PDF 的快速解析路径。
    page_contents: list[PdfPageContent] = []
    document = fitz.open(str(path))
    try:
        page_count = document.page_count
        for page_index in range(page_count):
            page_number = page_index + 1
            page = document.load_page(page_index)
            raw_text = page.get_text("text") or ""
            text, table_like_lines = _normalize_pdf_page_text(raw_text)
            page_contents.append(
                PdfPageContent(
                    page_number=page_number,
                    raw_text=raw_text,
                    text=text,
                    table_like_lines=table_like_lines,
                )
            )
    finally:
        document.close()

    text_page_reports = [
        _pdf_page_report(
            page_number=content.page_number,
            raw_text=content.raw_text,
            clean_text_value=content.text,
            table_like_lines=content.table_like_lines,
            extraction_method=content.extraction_method,
            ocr_confidence=content.ocr_confidence,
        )
        for content in page_contents
    ]
    text_quality = _pdf_quality_summary(
        page_reports=text_page_reports,
        page_count=page_count,
    )
    ocr_candidate_pages = _pdf_pages_requiring_ocr(text_page_reports)
    ocr_applied = False
    ocr_unavailable_reason: str | None = None
    ocr_results: list[PdfOcrPageResult] = []
    if ocr_candidate_pages and _pdf_ocr_enabled():
        try:
            ocr_results = _ocr_pdf_pages(path, ocr_candidate_pages)
            page_contents = _merge_pdf_ocr_results(page_contents, ocr_results)
            ocr_applied = any(result.text.strip() for result in ocr_results)
        except RuntimeError as exc:
            ocr_unavailable_reason = str(exc)
            if text_quality["ocr_required"]:
                raise RuntimeError(
                    f"PDF 疑似扫描件或低文本质量文档，文字解析不可用且 OCR 兜底失败: {path}。{exc}"
                ) from exc

    page_reports = [
        _pdf_page_report(
            page_number=content.page_number,
            raw_text=content.raw_text,
            clean_text_value=content.text,
            table_like_lines=content.table_like_lines,
            extraction_method=content.extraction_method,
            ocr_confidence=content.ocr_confidence,
        )
        for content in page_contents
    ]
    quality = _pdf_quality_summary(page_reports=page_reports, page_count=page_count)
    warnings = list(quality["warnings"])
    if ocr_candidate_pages and not _pdf_ocr_enabled():
        warnings = _dedupe([*warnings, "ocr_disabled"])
    if ocr_unavailable_reason:
        warnings = _dedupe([*warnings, "ocr_fallback_unavailable"])

    lines = [f"# {path.stem}"]
    extracted_pages = 0
    for content in page_contents:
        if not content.text:
            continue
        extracted_pages += 1
        lines.extend(["", f"## 第{content.page_number}页", content.text])

    cleaned = clean_text("\n".join(lines))
    if cleaned == f"# {path.stem}":
        raise ValueError(
            f"PDF 未抽取到可用文本，疑似扫描件或图片型 PDF，需要安装并启用 OCR 后再入库: {path}"
        )
    ocr_pages = [
        content.page_number
        for content in page_contents
        if content.extraction_method == "ocr"
    ]
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": "pdf",
            "title": path.stem,
            "media_type": "pdf",
            "parser_strategy": "pymupdf_text_with_ocr_fallback" if ocr_applied else "pymupdf_text_with_page_quality_report",
            "pdf_document_class": _classify_pdf_document(page_reports, page_count),
            "page_count": page_count,
            "extracted_page_count": extracted_pages,
            "empty_page_count": quality["empty_page_count"],
            "low_text_page_count": quality["low_text_page_count"],
            "scanned_page_candidates": quality["scanned_page_candidates"],
            "table_like_page_count": quality["table_like_page_count"],
            "parse_quality_status": quality["status"],
            "parse_warnings": warnings,
            "ocr_required": quality["ocr_required"],
            "ocr_candidate_pages": ocr_candidate_pages,
            "ocr_applied": ocr_applied,
            "ocr_provider": "pymupdf+pytesseract" if ocr_applied else None,
            "ocr_page_count": len(ocr_pages),
            "ocr_pages": ocr_pages,
            "ocr_confidence_avg": _average_ocr_confidence(page_contents),
            "ocr_unavailable_reason": ocr_unavailable_reason,
            "pdf_page_reports": [report.__dict__ for report in page_reports],
        },
    )


def _parse_pptx(path: Path) -> ParsedFile | None:
    try:
        from pptx import Presentation
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "解析 PPTX 需要安装可选依赖：pip install -e \".[multimodal]\""
        ) from exc

    # PPTX 按幻灯片遍历文本 shape。
    # 每页幻灯片被转成独立小节，后续可通过 slide_number 追溯来源。
    presentation = Presentation(str(path))
    lines = [f"# {path.stem}"]
    extracted_slides = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            content = getattr(shape, "text", None)
            if content is None:
                continue
            content = clean_text(str(content))
            if content:
                texts.append(content)
        if not texts:
            continue
        extracted_slides += 1
        lines.extend(
            [
                "",
                f"## 第{slide_number}页幻灯片",
                "\n".join(texts),
            ]
        )
    cleaned = clean_text("\n".join(lines))
    if cleaned == f"# {path.stem}":
        return None
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": "pptx",
            "title": path.stem,
            "media_type": "pptx",
            "slide_count": len(presentation.slides),
            "extracted_slide_count": extracted_slides,
        },
    )


class _HtmlToMarkdownParser(HTMLParser):
    """把 HTML DOM 中的标题、段落、列表和表格转为稳定 Markdown。"""

    def __init__(self, default_title: str) -> None:
        super().__init__(convert_charrefs=True)
        self.default_title = default_title
        self.title: str | None = None
        self.link_count = 0
        self.table_count = 0
        self._lines: list[str] = []
        self._current_tag: str | None = None
        self._current_heading_level: int | None = None
        self._current_text: list[str] = []
        self._current_cell: list[str] | None = None
        self._current_row: list[str] | None = None
        self._table_headers: list[str] | None = None
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        tag = tag.lower()
        if tag == "a":
            self.link_count += 1
        if tag == "title":
            self._in_title = True
            self._current_text = []
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._current_tag = tag
            self._current_heading_level = int(tag[1])
            self._current_text = []
            return
        if tag in {"p", "li"}:
            self._current_tag = tag
            self._current_text = []
            return
        if tag == "table":
            self.table_count += 1
            self._table_headers = None
            self._append_line("")
            self._append_line(f"## 表格 {self.table_count}")
            return
        if tag == "tr":
            self._current_row = []
            return
        if tag in {"td", "th"}:
            self._current_cell = []
            return
        if tag == "br":
            self._append_text("\n")

    def handle_data(self, data: str) -> None:
        self._append_text(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "title":
            title = clean_text(" ".join(self._current_text))
            if title:
                self.title = title
            self._in_title = False
            self._current_text = []
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"} and self._current_tag == tag:
            text = clean_text(" ".join(self._current_text))
            if text:
                level = min(self._current_heading_level or 1, 6)
                self._append_line("")
                self._append_line(f"{'#' * level} {text}")
            self._reset_current_text()
            return
        if tag in {"p", "li"} and self._current_tag == tag:
            text = clean_text(" ".join(self._current_text))
            if text:
                prefix = "- " if tag == "li" else ""
                self._append_line(f"{prefix}{text}")
            self._reset_current_text()
            return
        if tag in {"td", "th"} and self._current_cell is not None:
            cell = clean_text(" ".join(self._current_cell))
            if self._current_row is not None:
                self._current_row.append(cell)
            self._current_cell = None
            return
        if tag == "tr" and self._current_row is not None:
            lines = _table_rows_to_lines([self._current_row], self._table_headers)
            if self._table_headers is None and self._current_row:
                self._table_headers = [cell for cell in self._current_row if cell]
            self._lines.extend(lines)
            self._current_row = None

    def markdown_lines(self) -> list[str]:
        lines = list(self._lines)
        title = self.title or self.default_title
        if not lines or not lines[0].startswith("# "):
            lines.insert(0, f"# {title}")
        return lines

    def _append_text(self, text: str) -> None:
        if self._current_cell is not None:
            self._current_cell.append(text)
        elif self._current_tag is not None or self._in_title:
            self._current_text.append(text)

    def _append_line(self, line: str) -> None:
        if line or (self._lines and self._lines[-1]):
            self._lines.append(line)

    def _reset_current_text(self) -> None:
        self._current_tag = None
        self._current_heading_level = None
        self._current_text = []


def _docx_heading_level(style_name: str) -> int | None:
    match = re.search(r"heading\s*(\d+)", style_name)
    if match:
        return int(match.group(1))
    if "标题" in style_name:
        match = re.search(r"(\d+)", style_name)
        return int(match.group(1)) if match else 1
    return None


def _table_rows_to_lines(
    rows: list[list[str]],
    headers: list[str] | None = None,
) -> list[str]:
    normalized_rows = [
        [cell.strip() for cell in row]
        for row in rows
        if any(str(cell).strip() for cell in row)
    ]
    if not normalized_rows:
        return []

    local_headers = headers
    lines: list[str] = []
    start_index = 0
    if local_headers is None:
        local_headers = normalized_rows[0]
        start_index = 1
        lines.append(f"表格列：{'；'.join(local_headers)}")
    for row in normalized_rows[start_index:]:
        if local_headers and len(row) == len(local_headers):
            pairs = [
                f"{header}={value}"
                for header, value in zip(local_headers, row)
                if header or value
            ]
            lines.append(f"表格行：{'；'.join(pairs)}")
        else:
            lines.append(f"表格行：{'；'.join(row)}")
    return lines


def _stringify_excel_cell(value: Any) -> str:
    if value is None:
        return ""
    return clean_text(str(value))


def _parse_srt(path: Path) -> ParsedFile | None:
    text = _read_text_with_fallback(path)
    cues = _parse_srt_cues(text)
    return _subtitle_cues_to_document(path, cues, parser_name="srt")


def _parse_vtt(path: Path) -> ParsedFile | None:
    text = _read_text_with_fallback(path)
    cues = _parse_vtt_cues(text)
    return _subtitle_cues_to_document(path, cues, parser_name="vtt")


def _normalize_pdf_page_text(raw_text: str) -> tuple[str, int]:
    """清洗单页 PDF 文本，并尽量保留表格行语义。"""

    cleaned = clean_text(raw_text)
    if not cleaned:
        return "", 0

    normalized_lines: list[str] = []
    table_headers: list[str] | None = None
    table_like_lines = 0
    for line in cleaned.splitlines():
        normalized, table_headers, is_table_like = _normalize_pdf_table_line(
            line,
            table_headers,
        )
        if is_table_like:
            table_like_lines += 1
        normalized_lines.append(normalized)
    return clean_text("\n".join(normalized_lines)), table_like_lines


def _normalize_pdf_table_line(
    line: str,
    table_headers: list[str] | None,
) -> tuple[str, list[str] | None, bool]:
    """把疑似 PDF 表格行转成行级文本，避免列关系完全丢失。"""

    stripped = line.strip()
    cells = [
        cell.strip()
        for cell in PDF_TABLE_CELL_SPLIT_RE.split(stripped)
        if cell.strip()
    ]
    if len(cells) < 2:
        return line, None, False

    if table_headers is None:
        return f"表格列：{'；'.join(cells)}", cells, True

    if len(cells) == len(table_headers):
        pairs = [
            f"{header}={value}"
            for header, value in zip(table_headers, cells)
        ]
        return f"表格行：{'；'.join(pairs)}", table_headers, True

    return f"表格行：{'；'.join(cells)}", table_headers, True


def _pdf_page_report(
    page_number: int,
    raw_text: str,
    clean_text_value: str,
    table_like_lines: int,
    extraction_method: str = "text",
    ocr_confidence: float | None = None,
) -> PdfPageReport:
    warning_codes: list[str] = []
    clean_length = len(clean_text_value)
    if clean_length == 0:
        status = "empty"
        warning_codes.append("empty_page")
    elif clean_length < PDF_MIN_TEXT_CHARS_PER_PAGE:
        status = "low_text"
        warning_codes.append("low_text_page")
    else:
        status = "parsed"

    garbled_ratio = _garbled_ratio(clean_text_value)
    if garbled_ratio > PDF_GARBLED_RATIO_WARNING:
        warning_codes.append("high_garbled_ratio")
        if status == "parsed":
            status = "parsed_with_warning"

    if table_like_lines > 0:
        warning_codes.append("table_like_text_detected")
    if extraction_method == "ocr":
        warning_codes.append("ocr_applied")
        if ocr_confidence is not None and ocr_confidence < PDF_OCR_CONFIDENCE_WARNING:
            warning_codes.append("low_ocr_confidence")
            if status == "parsed":
                status = "parsed_with_warning"

    return PdfPageReport(
        page_number=page_number,
        raw_text_length=len(raw_text or ""),
        clean_text_length=clean_length,
        status=status,
        warning_codes=warning_codes,
        table_like_lines=table_like_lines,
        extraction_method=extraction_method,
        ocr_confidence=ocr_confidence,
    )


def _pdf_quality_summary(
    page_reports: list[PdfPageReport],
    page_count: int,
) -> dict[str, Any]:
    empty_page_count = sum(1 for report in page_reports if report.status == "empty")
    low_text_page_count = sum(1 for report in page_reports if report.status == "low_text")
    table_like_page_count = sum(
        1
        for report in page_reports
        if report.table_like_lines > 0
    )
    scanned_page_candidates = empty_page_count + low_text_page_count
    warnings = _dedupe(
        [
            warning
            for report in page_reports
            for warning in report.warning_codes
        ]
    )
    if page_count > 0 and empty_page_count == page_count:
        status = "parse_failed"
    elif page_count > 0 and scanned_page_candidates / page_count >= 0.5:
        status = "needs_review"
    elif warnings:
        status = "parsed_with_warning"
    else:
        status = "parsed_success"
    return {
        "status": status,
        "warnings": warnings,
        "empty_page_count": empty_page_count,
        "low_text_page_count": low_text_page_count,
        "scanned_page_candidates": scanned_page_candidates,
        "table_like_page_count": table_like_page_count,
        "ocr_required": status in {"parse_failed", "needs_review"},
    }


def _pdf_pages_requiring_ocr(page_reports: list[PdfPageReport]) -> list[int]:
    """识别需要 OCR 兜底的 PDF 页码。"""

    pages: list[int] = []
    for report in page_reports:
        if report.extraction_method == "ocr":
            continue
        if report.status in {"empty", "low_text"}:
            pages.append(report.page_number)
            continue
        if "high_garbled_ratio" in report.warning_codes:
            pages.append(report.page_number)
    return pages


def _pdf_ocr_enabled() -> bool:
    return _env_bool("RAG_PDF_OCR_ENABLED", True)


def _ocr_pdf_pages(path: Path, page_numbers: list[int]) -> list[PdfOcrPageResult]:
    """用 OCR 解析指定 PDF 页面。

    实现选择 `PyMuPDF + pytesseract`，前者负责把 PDF 页渲染成图像，
    后者负责文字识别；OCR 只有在 PDF 文字抽取不足时才会用到。
    """

    if not page_numbers:
        return []
    try:
        import fitz
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "OCR PDF 需要安装 PyMuPDF：pip install -e \".[ocr]\"。"
        ) from exc
    try:
        import pytesseract
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "OCR PDF 需要安装可选依赖：pip install -e \".[ocr]\"，缺少 pytesseract。"
        ) from exc
    try:
        from PIL import Image
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "OCR PDF 需要安装可选依赖：pip install -e \".[ocr]\"，缺少 Pillow。"
        ) from exc

    tesseract_cmd = os.getenv("RAG_TESSERACT_CMD")
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    dpi = _env_int("RAG_PDF_OCR_DPI", 200)
    scale = dpi / 72
    language = os.getenv("RAG_PDF_OCR_LANG", "chi_sim+eng")
    target_pages = set(page_numbers)
    results: list[PdfOcrPageResult] = []
    document = fitz.open(str(path))
    try:
        for page_index in range(document.page_count):
            page_number = page_index + 1
            if page_number not in target_pages:
                continue
            page = document.load_page(page_index)
            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(scale, scale),
                alpha=False,
            )
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            try:
                text, confidence = _run_tesseract_ocr(
                    pytesseract=pytesseract,
                    image=image,
                    language=language,
                )
            finally:
                image.close()
            results.append(
                PdfOcrPageResult(
                    page_number=page_number,
                    text=text,
                    confidence=confidence,
                )
            )
    finally:
        document.close()
    return results


def _run_tesseract_ocr(pytesseract, image, language: str) -> tuple[str, float | None]:
    try:
        data = pytesseract.image_to_data(
            image,
            lang=language,
            output_type=pytesseract.Output.DICT,
        )
    except Exception as exc:  # noqa: BLE001 - OCR 依赖外部二进制和语言包，需保留原始错误。
        raise RuntimeError(
            "Tesseract OCR 执行失败，请确认已安装 tesseract 二进制、语言包，"
            "并按需配置 RAG_TESSERACT_CMD / RAG_PDF_OCR_LANG。"
        ) from exc

    tokens: list[str] = []
    confidences: list[float] = []
    texts = data.get("text", [])
    raw_confidences = data.get("conf", [])
    for token, confidence in zip(texts, raw_confidences):
        token = str(token).strip()
        if token:
            tokens.append(token)
        try:
            confidence_value = float(confidence)
        except (TypeError, ValueError):
            continue
        if confidence_value >= 0:
            confidences.append(confidence_value / 100)
    text = clean_text(" ".join(tokens))
    average_confidence = (
        sum(confidences) / len(confidences)
        if confidences
        else None
    )
    return text, average_confidence


def _merge_pdf_ocr_results(
    page_contents: list[PdfPageContent],
    ocr_results: list[PdfOcrPageResult],
) -> list[PdfPageContent]:
    """把 OCR 结果合并回 PDF 页内容，优先替换低质量文字抽取结果。"""

    ocr_by_page = {result.page_number: result for result in ocr_results}
    merged: list[PdfPageContent] = []
    for content in page_contents:
        ocr_result = ocr_by_page.get(content.page_number)
        if ocr_result is None or not ocr_result.text.strip():
            merged.append(content)
            continue
        ocr_text, table_like_lines = _normalize_pdf_page_text(ocr_result.text)
        if not _should_replace_pdf_text_with_ocr(content, ocr_text):
            merged.append(content)
            continue
        merged.append(
            PdfPageContent(
                page_number=content.page_number,
                raw_text=content.raw_text,
                text=ocr_text,
                table_like_lines=table_like_lines,
                extraction_method="ocr",
                ocr_confidence=ocr_result.confidence,
            )
        )
    return merged


def _should_replace_pdf_text_with_ocr(
    content: PdfPageContent,
    ocr_text: str,
) -> bool:
    if not ocr_text.strip():
        return False
    if not content.text.strip():
        return True
    if len(content.text) < PDF_MIN_TEXT_CHARS_PER_PAGE:
        return True
    if _garbled_ratio(content.text) > PDF_GARBLED_RATIO_WARNING:
        return True
    return len(ocr_text) > len(content.text) * 1.5


def _classify_pdf_document(
    page_reports: list[PdfPageReport],
    page_count: int,
) -> str:
    if page_count <= 0:
        return "empty_pdf"
    ocr_pages = sum(1 for report in page_reports if report.extraction_method == "ocr")
    scanned_candidates = sum(
        1
        for report in page_reports
        if report.status in {"empty", "low_text"}
    )
    table_like_pages = sum(1 for report in page_reports if report.table_like_lines > 0)
    if ocr_pages == page_count:
        return "ocr_pdf"
    if ocr_pages > 0:
        return "mixed_text_ocr_pdf"
    if scanned_candidates == page_count:
        return "scanned_pdf"
    if scanned_candidates > 0:
        return "mixed_text_scanned_pdf"
    if table_like_pages > 0:
        return "table_or_layout_pdf"
    return "text_pdf"


def _average_ocr_confidence(
    page_contents: list[PdfPageContent],
) -> float | None:
    confidences = [
        content.ocr_confidence
        for content in page_contents
        if content.extraction_method == "ocr" and content.ocr_confidence is not None
    ]
    if not confidences:
        return None
    return sum(confidences) / len(confidences)


def _garbled_ratio(text: str) -> float:
    if not text:
        return 0.0
    suspicious = sum(
        1
        for char in text
        if char == "\ufffd" or unicodedata.category(char) == "Cc"
    )
    return suspicious / len(text)


def _dedupe(items: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for item in items:
        if item in seen:
            continue
        seen.add(item)
        result.append(item)
    return result


def _document_version_id(content_hash: str) -> str:
    return f"v_{content_hash[:16]}"


def _normalize_tenant_id(metadata: dict[str, Any]) -> str:
    value = (
        metadata.get("tenant_id")
        or metadata.get("tenant")
        or metadata.get("project_id")
        or "default"
    )
    return str(value).strip() or "default"


def _normalize_permission_tags(metadata: dict[str, Any]) -> list[str]:
    raw = (
        metadata.get("permission_tags")
        or metadata.get("permissions")
        or metadata.get("permission")
        or "public"
    )
    if isinstance(raw, str):
        values = re.split(r"[,，;；\s]+", raw)
    elif isinstance(raw, (list, tuple, set)):
        values = [str(item) for item in raw]
    else:
        values = [str(raw)]
    tags = [value.strip() for value in values if value and value.strip()]
    return _dedupe(tags) or ["public"]


def _build_parsed_blocks(
    document_id: str,
    version_id: str,
    text: str,
    document_metadata: dict[str, Any],
) -> list[ParsedBlock]:
    lines = text.splitlines()
    block_specs: list[dict[str, Any]] = []
    heading_stack: list[tuple[int, str, str]] = []
    body_lines: list[str] = []

    def current_section_path() -> list[str]:
        return [title for _, title, _ in heading_stack]

    def current_parent_id() -> str | None:
        return heading_stack[-1][2] if heading_stack else None

    def append_body_block() -> None:
        body = clean_text("\n".join(body_lines))
        body_lines.clear()
        if not body:
            return
        section_path = current_section_path()
        location = _extract_block_location(section_path)
        block_specs.append(
            {
                "block_type": _infer_block_type(body),
                "text": body,
                "section_path": section_path,
                "parent_block_id": current_parent_id(),
                "page_number": location.get("page_number"),
                "metadata": {
                    **_common_block_metadata(document_metadata),
                    **location,
                    **_error_code_block_metadata(
                        body,
                        section_path,
                        document_metadata,
                    ),
                },
            }
        )

    for line in lines:
        match = HEADING_RE.match(line)
        if not match:
            body_lines.append(line)
            continue

        append_body_block()
        level = len(match.group(1))
        title = match.group(2).strip()
        heading_stack = [
            (existing_level, existing_title, existing_block_id)
            for existing_level, existing_title, existing_block_id in heading_stack
            if existing_level < level
        ]
        heading_block_id = _parsed_block_id(document_id, len(block_specs), title)
        section_path = [title for _, title, _ in heading_stack] + [title]
        location = _extract_block_location(section_path)
        block_specs.append(
            {
                "block_id": heading_block_id,
                "block_type": "heading",
                "text": title,
                "section_path": section_path,
                "parent_block_id": heading_stack[-1][2] if heading_stack else None,
                "page_number": location.get("page_number"),
                "metadata": {
                    **_common_block_metadata(document_metadata),
                    "heading_level": level,
                    **location,
                    **_error_code_block_metadata(
                        title,
                        section_path,
                        document_metadata,
                    ),
                },
            }
        )
        heading_stack.append((level, title, heading_block_id))

    append_body_block()
    if not block_specs and text.strip():
        body = clean_text(text)
        block_specs.append(
            {
                "block_type": _infer_block_type(body),
                "text": body,
                "section_path": [],
                "parent_block_id": None,
                "page_number": None,
                "metadata": _common_block_metadata(document_metadata),
            }
        )

    blocks: list[ParsedBlock] = []
    for index, spec in enumerate(block_specs):
        block_id = spec.get("block_id") or _parsed_block_id(
            document_id,
            index,
            str(spec["text"]),
        )
        blocks.append(
            ParsedBlock(
                block_id=block_id,
                document_id=document_id,
                version_id=version_id,
                block_type=str(spec["block_type"]),
                text=str(spec["text"]),
                raw_text=str(spec["text"]),
                clean_text=str(spec["text"]),
                page_number=spec.get("page_number"),
                section_path=list(spec.get("section_path") or []),
                parent_block_id=spec.get("parent_block_id"),
                previous_block_id=blocks[-1].block_id if blocks else None,
                metadata=dict(spec.get("metadata") or {}),
                quality_status=str(
                    document_metadata.get("parse_quality_status", "parsed_success")
                ),
            )
        )

    # dataclass 冻结后无法原地补 next_block_id，这里重建一次，形成可导航的文档树链路。
    linked_blocks: list[ParsedBlock] = []
    for index, block in enumerate(blocks):
        linked_blocks.append(
            ParsedBlock(
                block_id=block.block_id,
                document_id=block.document_id,
                version_id=block.version_id,
                block_type=block.block_type,
                text=block.text,
                raw_text=block.raw_text,
                clean_text=block.clean_text,
                page_number=block.page_number,
                bbox=block.bbox,
                section_path=block.section_path,
                parent_block_id=block.parent_block_id,
                previous_block_id=block.previous_block_id,
                next_block_id=blocks[index + 1].block_id if index + 1 < len(blocks) else None,
                metadata=block.metadata,
                quality_status=block.quality_status,
            )
        )
    return linked_blocks


def _common_block_metadata(document_metadata: dict[str, Any]) -> dict[str, Any]:
    return {
        "source": document_metadata.get("source"),
        "source_path": document_metadata.get("source_path"),
        "filename": document_metadata.get("filename"),
        "title": document_metadata.get("title"),
        "parser_name": document_metadata.get("parser_name"),
        "parser_version": document_metadata.get("parser_version"),
        "tenant_id": document_metadata.get("tenant_id"),
        "permission_tags": document_metadata.get("permission_tags", ["public"]),
        "source_type": document_metadata.get("source_type"),
        "business_module": document_metadata.get("business_module"),
    }


def _error_code_block_metadata(
    text: str,
    section_path: list[str],
    document_metadata: dict[str, Any],
) -> dict[str, Any]:
    local_text = " ".join([*section_path, text])
    codes = extract_error_codes(local_text)
    if not codes:
        document_codes = _normalize_error_code_list(document_metadata.get("error_codes"))
        if len(document_codes) == 1:
            codes = document_codes
    metadata: dict[str, Any] = {"error_codes": codes}
    if not codes:
        return metadata
    metadata["error_code"] = codes[0]
    metadata["intent_labels"] = ["explain_error"]
    metadata["keywords"] = codes
    return metadata


def _normalize_error_code_list(value: Any) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, str):
        raw_values = re.split(r"[,，;；\s]+", value)
    elif isinstance(value, (list, tuple, set)):
        raw_values = [str(item) for item in value]
    else:
        raw_values = [str(value)]
    result: list[str] = []
    seen: set[str] = set()
    for item in raw_values:
        if not item or not item.strip():
            continue
        code = normalize_error_code(item)
        if code in seen:
            continue
        seen.add(code)
        result.append(code)
    return result


def _extract_block_location(section_path: list[str]) -> dict[str, str | int]:
    for title in reversed(section_path):
        page_match = PAGE_TITLE_RE.match(title)
        if page_match:
            return {"page_number": int(page_match.group(1))}
        slide_match = SLIDE_TITLE_RE.match(title)
        if slide_match:
            return {"slide_number": int(slide_match.group(1))}
        segment_match = SEGMENT_TITLE_RE.match(title)
        if segment_match:
            return {"time_range": segment_match.group(1)}
    return {}


def _infer_block_type(text: str) -> str:
    stripped = text.strip()
    if stripped.startswith("表格列：") or stripped.startswith("表格行："):
        return "table"
    if stripped.startswith("- "):
        return "list"
    if stripped.startswith("```"):
        return "code"
    return "paragraph"


def _parsed_block_id(document_id: str, index: int, text: str) -> str:
    digest = hashlib.sha256()
    digest.update(document_id.encode("utf-8"))
    digest.update(str(index).encode("utf-8"))
    digest.update(text.encode("utf-8"))
    return f"blk_{digest.hexdigest()[:20]}"


def _subtitle_cues_to_document(
    path: Path,
    cues: list[SubtitleCue],
    parser_name: str,
) -> ParsedFile | None:
    if not cues:
        return None
    # 字幕文件按时间轴切成片段，小节标题保留 start-end，方便召回后定位视频时间段。
    lines = [f"# {path.stem}"]
    for index, cue in enumerate(cues, start=1):
        lines.extend(
            [
                "",
                f"## 片段 {index} [{cue.start}-{cue.end}]",
                cue.text,
            ]
        )
    cleaned = clean_text("\n".join(lines))
    if cleaned == f"# {path.stem}":
        return None
    return ParsedFile(
        text=cleaned,
        metadata={
            "parser_name": parser_name,
            "title": path.stem,
            "media_type": "video_transcript",
            "segment_count": len(cues),
        },
    )


def _parse_srt_cues(text: str) -> list[SubtitleCue]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = [block.strip() for block in re.split(r"\n{2,}", normalized) if block.strip()]
    cues: list[SubtitleCue] = []
    for block in blocks:
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        if not lines:
            continue
        if "-->" in lines[0]:
            time_line_index = 0
        elif len(lines) > 1 and "-->" in lines[1]:
            time_line_index = 1
        else:
            continue
        start, end = _split_time_range(lines[time_line_index])
        if not start or not end:
            continue
        content_lines = lines[time_line_index + 1 :]
        content = clean_text(" ".join(content_lines))
        if not content:
            continue
        cues.append(SubtitleCue(start=start, end=end, text=content))
    return cues


def _parse_vtt_cues(text: str) -> list[SubtitleCue]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line for line in normalized.split("\n")]
    if lines and lines[0].strip().upper().startswith("WEBVTT"):
        lines = lines[1:]
    blocks = [block.strip() for block in re.split(r"\n{2,}", "\n".join(lines)) if block.strip()]
    cues: list[SubtitleCue] = []
    for block in blocks:
        cue_lines = [line.strip() for line in block.split("\n") if line.strip()]
        if not cue_lines:
            continue
        if "-->" in cue_lines[0]:
            time_line = cue_lines[0]
            content_lines = cue_lines[1:]
        elif len(cue_lines) > 1 and "-->" in cue_lines[1]:
            time_line = cue_lines[1]
            content_lines = cue_lines[2:]
        else:
            continue
        start, end = _split_time_range(time_line)
        content = clean_text(" ".join(content_lines))
        if not start or not end or not content:
            continue
        cues.append(SubtitleCue(start=start, end=end, text=content))
    return cues


def _split_time_range(time_line: str) -> tuple[str, str]:
    left, _, right = time_line.partition("-->")
    if not _:
        return "", ""
    start = left.strip()
    # 移除 VTT 时轴行中的附加样式信息。
    end = right.strip().split()[0]
    return start, end


def _read_text_with_fallback(path: Path) -> str:
    # 业务资料经常来自不同系统，编码不稳定；这里按常见中文场景做 fallback。
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _env_bool(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None or value.strip() == "":
        return default
    normalized = value.strip().lower()
    if normalized in {"1", "true", "yes", "y", "on"}:
        return True
    if normalized in {"0", "false", "no", "n", "off"}:
        return False
    raise ValueError(f"{name} must be a boolean, got {value!r}")


def _env_int(name: str, default: int) -> int:
    value = os.getenv(name)
    if value is None or value.strip() == "":
        return default
    try:
        parsed = int(value)
    except ValueError as exc:
        raise ValueError(f"{name} must be an integer, got {value!r}") from exc
    if parsed <= 0:
        raise ValueError(f"{name} must be greater than 0")
    return parsed


def _extract_title(text: str) -> str | None:
    match = TITLE_RE.search(text)
    if match:
        return match.group(1).strip()
    return None


def _extract_front_matter(text: str) -> tuple[dict[str, str], str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    match = FRONT_MATTER_RE.match(normalized)
    if not match:
        return {}, text

    metadata: dict[str, str] = {}
    for line in match.group(1).splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        key, separator, value = line.partition(":")
        if not separator:
            continue
        key = key.strip()
        value = value.strip().strip("'\"")
        if key and value:
            metadata[key] = value
    return metadata, normalized[match.end():]


def _json_records(path: Path, data) -> tuple[str, list]:
    if isinstance(data, list):
        return path.stem, data
    if isinstance(data, dict):
        title = str(data.get("title") or data.get("name") or path.stem)
        for key in JSON_RECORD_KEYS:
            value = data.get(key)
            if isinstance(value, list):
                return title, value
        return title, [data]
    return path.stem, [data]


def _json_record_title(index: int, record) -> str:
    if isinstance(record, dict):
        for key in ("title", "name", "question", "error_code", "code", "order_id", "id"):
            value = record.get(key)
            if value not in ("", None):
                return str(value)
    return f"记录 {index}"


def _stringify_json_value(value) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def _infer_source_type(path: Path, text: str) -> str:
    target = f"{path.as_posix()} {text[:800]}".lower()
    if path.suffix.lower() in {".srt", ".vtt"}:
        return "video_transcript"
    if path.suffix.lower() == ".pptx":
        return "training_material"
    if path.suffix.lower() == ".pdf":
        return "manual"
    if _contains_any(target, ("故障", "案例", "fault", "case")):
        return "fault_case"
    if _contains_any(target, ("异常码", "错误码", "error_code", "error-code")):
        return "error_code"
    if _contains_any(target, ("接口", "api", "返回说明")):
        return "api_doc"
    if _contains_any(target, ("派单", "dispatch")):
        return "dispatch_rule"
    if _contains_any(target, ("地址", "address")):
        return "address_rule"
    if _contains_any(target, ("设备", "光猫", "onu", "ont", "device")):
        return "device_rule"
    if _contains_any(target, ("手册", "manual", "guide")):
        return "manual"
    return "manual"


def _infer_business_module(path: Path, text: str) -> str:
    target = f"{path.as_posix()} {text[:800]}".lower()
    if path.suffix.lower() in {".srt", ".vtt", ".pptx"}:
        return "培训资料"
    if _contains_any(target, ("故障", "fault", "case")):
        return "故障处理"
    if _contains_any(target, ("地址", "address")):
        return "地址校验"
    if _contains_any(target, ("派单", "dispatch")):
        return "智能派单"
    if _contains_any(target, ("设备", "光猫", "onu", "ont", "pon", "los", "device")):
        return "设备维护"
    if _contains_any(target, ("异常码", "错误码", "接口", "api", "error")):
        return "异常处理"
    if _contains_any(target, ("工单", "work_order", "order")):
        return "工单流转"
    return "通用知识"


def _contains_any(text: str, keywords: tuple[str, ...]) -> bool:
    return any(keyword.lower() in text for keyword in keywords)


def _relative_source(path: Path, source_dir: Path | None) -> str:
    if source_dir is None:
        return str(path)
    try:
        return str(path.relative_to(source_dir))
    except ValueError:
        return str(path)


def _document_id(source_path: str, content_hash: str) -> str:
    digest = hashlib.sha256()
    digest.update(source_path.encode("utf-8"))
    digest.update(b"\0")
    digest.update(content_hash.encode("utf-8"))
    digest.update(b"\0")
    digest.update(PARSER_VERSION.encode("utf-8"))
    return digest.hexdigest()[:24]
